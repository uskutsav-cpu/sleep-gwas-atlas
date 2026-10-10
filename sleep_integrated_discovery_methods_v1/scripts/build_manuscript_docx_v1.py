"""Format the reviewed manuscript text as an editable document, without analysis."""
import argparse
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def add_link(paragraph, label, url):
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '1F497D'); props.append(color)
    run.append(props)
    text = OxmlElement('w:t'); text.text = label; run.append(text)
    link.append(run); paragraph._p.append(link)


def inline(paragraph, text):
    # Keep references clickable and preserve emphasis without exposing Markdown.
    pattern = r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|r_g|10\^−15)'
    for token in re.split(pattern, text):
        if not token:
            continue
        match = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', token)
        if match:
            add_link(paragraph, match[1], match[2])
        elif token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith('*') and token.endswith('*'):
            paragraph.add_run(token[1:-1]).italic = True
        elif token.startswith('`') and token.endswith('`'):
            paragraph.add_run(token[1:-1])
        elif token == 'r_g':
            paragraph.add_run('r')
            paragraph.add_run('g').font.subscript = True
        elif token == '10^−15':
            paragraph.add_run('10')
            paragraph.add_run('−15').font.superscript = True
        else:
            paragraph.add_run(token)


def heading_text(text):
    # Titles and headings use words and numbers, with punctuation in body text.
    return re.sub(r'[^\w\s]', ' ', text, flags=re.UNICODE).strip()


def build(source, output):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = doc.styles['Normal']
    normal.font.name = 'Times New Roman'; normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size in [('Title', 17), ('Heading 1', 13), ('Heading 2', 11.5), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        # Remove inherited decorative title/heading rules from the Word template.
        for border in list(style._element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
        style.font.bold = name != 'Title'
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.core_properties.title = 'Sleep genetics manuscript for human scientific review'
    doc.core_properties.author = ''
    doc.core_properties.subject = 'One integrated qualified sleep genetics paper'
    doc.core_properties.created = datetime.now(timezone.utc)
    doc.core_properties.modified = datetime.now(timezone.utc)
    lines = source.read_text().splitlines()
    paragraph_lines = []

    def flush():
        if paragraph_lines:
            p = doc.add_paragraph()
            inline(p, ' '.join(paragraph_lines))
            paragraph_lines.clear()

    for line in lines:
        if not line.strip():
            flush(); continue
        h = re.match(r'^(#{1,4})\s+(.+)$', line)
        if h:
            flush()
            if h[2].startswith('Supplementary display and source-data plan'):
                doc.add_page_break()
            level = len(h[1])
            p = doc.add_paragraph(style='Title' if level == 1 else f'Heading {min(level-1,3)}')
            p.add_run(heading_text(h[2]))
        elif line.startswith('- '):
            flush(); p = doc.add_paragraph(style='List Bullet'); inline(p, line[2:])
        elif re.match(r'^\d+\.\s', line):
            flush(); p = doc.add_paragraph(); inline(p, line)
        elif line.startswith('|'):
            raise ValueError('Manuscript contains a Markdown table; format it deliberately before delivery')
        elif line.strip() in ['---', '***']:
            flush()
        else:
            paragraph_lines.append(line)
    flush()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f'Saved {output}; {len(doc.paragraphs)} paragraphs. No calculations performed.')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source', type=Path)
    parser.add_argument('output', type=Path)
    args = parser.parse_args()
    build(args.source, args.output)
