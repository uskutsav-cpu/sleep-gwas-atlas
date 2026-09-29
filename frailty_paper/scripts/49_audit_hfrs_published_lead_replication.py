#!/usr/bin/env python3
"""Audit published HFRS lead-variant rows against the publisher workbook.

This is a descriptive reproduction of values reported in Mak et al. (2025).
It does not analyze sleep/frailty pairs, estimate genome-wide effects, or
resolve the exact FinnGen HFRS summary-statistics access gate.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_SOURCE_SHA256 = "123c340cf063e2ef618b8ef543dedea6b7c8ef69008488f01027eedbffc04625"
SOURCE_RELATIVE = Path(
    "frailty_paper/data/gwas/hfrs_mak_2025/mak_2025_hfrs_supplementary_tables.xlsx"
)
OUTPUT_RELATIVE = Path("frailty_paper/analysis")
ARTICLE = "Mak et al., Nature Aging 5:1589–1600 (2025), doi:10.1038/s43587-025-00925-y"
ARTICLE_REPORTED = {
    "ST1": {"discovery_rows": 1588, "ukb_available_hits": 1262, "ukb_nominal_hits": 688,
            "ukb_gws_hits": 73, "lead_rows": 53, "ukb_available_leads": 36,
            "ukb_nominal_leads": 14, "ukb_gws_leads": 2},
    "ST2": {"discovery_rows": 492, "ukb_available_hits": 435, "ukb_nominal_hits": 118,
            "ukb_gws_hits": 21, "lead_rows": 42, "ukb_available_leads": 26,
            "ukb_nominal_leads": 10, "ukb_gws_leads": 1},
}
FIELDS = [
    "table", "snp", "chr", "pos", "gene", "locus", "novel_snp",
    "finngen_alt", "finngen_ref", "finngen_beta", "finngen_se", "finngen_p",
    "ukb_alt", "ukb_ref", "ukb_beta", "ukb_se", "ukb_p",
    "ukb_result_available_beta_se_p", "ukb_nominal_p_lt_0_05", "ukb_gws_p_lt_5e_8",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def is_present(value: Any) -> bool:
    return value is not None and str(value).strip() not in {"", ".", "NA", "N/A"}


def number(value: Any) -> float | None:
    if not is_present(value):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def truthy(value: Any) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "y"}


def excel_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    return value


def audit_sheet(sheet: Any, name: str) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    iterator = sheet.iter_rows(values_only=True)
    title_row = next(iterator, None)
    header = next(iterator, None)
    if header is None:
        raise ValueError(f"{name} has no tabular header")
    indices = {str(column): index for index, column in enumerate(header) if column is not None}
    required = {
        "SNP", "chr", "pos", "gene", "locus", "indep_lead_snp", "novel_snp",
        "finngen_alt", "finngen_ref", "finngen_beta", "finngen_SE", "finngen_p",
        "ukb_alt", "ukb_ref", "ukb_beta", "ukb_SE", "ukb_p",
    }
    missing = sorted(required - indices.keys())
    if missing:
        raise ValueError(f"{name} missing required columns: {missing}")

    discovery_rows: list[tuple[Any, ...]] = []
    lead_rows: list[tuple[Any, ...]] = []
    for row in iterator:
        # Some reported signals have no rsID (the cell is "."), but remain
        # valid rows identified by chromosome/position and association fields.
        if row is None or not any(is_present(value) for value in row):
            continue
        discovery_rows.append(row)
        if truthy(row[indices["indep_lead_snp"]]):
            lead_rows.append(row)

    records: list[dict[str, Any]] = []
    available = nominal = genome_wide = 0
    missing_fields: dict[str, int] = {"ukb_beta": 0, "ukb_SE": 0, "ukb_p": 0}
    for row in lead_rows:
        beta = number(row[indices["ukb_beta"]])
        se = number(row[indices["ukb_SE"]])
        pvalue = number(row[indices["ukb_p"]])
        for column, value in (("ukb_beta", beta), ("ukb_SE", se), ("ukb_p", pvalue)):
            if value is None:
                missing_fields[column] += 1
        row_available = beta is not None and se is not None and pvalue is not None
        available += int(row_available)
        is_nominal = row_available and pvalue < 0.05
        is_gws = row_available and pvalue < 5e-8
        nominal += int(is_nominal)
        genome_wide += int(is_gws)
        records.append({
            "table": name,
            "snp": excel_value(row[indices["SNP"]]),
            "chr": excel_value(row[indices["chr"]]),
            "pos": excel_value(row[indices["pos"]]),
            "gene": excel_value(row[indices["gene"]]),
            "locus": excel_value(row[indices["locus"]]),
            "novel_snp": excel_value(row[indices["novel_snp"]]),
            "finngen_alt": excel_value(row[indices["finngen_alt"]]),
            "finngen_ref": excel_value(row[indices["finngen_ref"]]),
            "finngen_beta": excel_value(row[indices["finngen_beta"]]),
            "finngen_se": excel_value(row[indices["finngen_SE"]]),
            "finngen_p": excel_value(row[indices["finngen_p"]]),
            "ukb_alt": excel_value(row[indices["ukb_alt"]]),
            "ukb_ref": excel_value(row[indices["ukb_ref"]]),
            "ukb_beta": excel_value(row[indices["ukb_beta"]]),
            "ukb_se": excel_value(row[indices["ukb_SE"]]),
            "ukb_p": excel_value(row[indices["ukb_p"]]),
            "ukb_result_available_beta_se_p": row_available,
            "ukb_nominal_p_lt_0_05": bool(is_nominal),
            "ukb_gws_p_lt_5e_8": bool(is_gws),
        })

    def count_results(rows: list[tuple[Any, ...]]) -> dict[str, int]:
        counts = {"available_beta_se_p": 0, "nominal_p_lt_0_05": 0, "gws_p_lt_5e_8": 0}
        for row in rows:
            beta = number(row[indices["ukb_beta"]])
            se = number(row[indices["ukb_SE"]])
            pvalue = number(row[indices["ukb_p"]])
            available_row = beta is not None and se is not None and pvalue is not None
            counts["available_beta_se_p"] += int(available_row)
            counts["nominal_p_lt_0_05"] += int(available_row and pvalue < 0.05)
            counts["gws_p_lt_5e_8"] += int(available_row and pvalue < 5e-8)
        return counts

    summary = {
        "discovery_rows": len(discovery_rows),
        "all_hit_ukb": count_results(discovery_rows),
        "lead_rows": len(lead_rows),
        "lead_ukb": {"available_beta_se_p": available, "nominal_p_lt_0_05": nominal,
                     "gws_p_lt_5e_8": genome_wide},
        "lead_rows_missing_ukb_field": missing_fields,
    }
    return records, summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repo", type=Path, default=Path("."))
    args = parser.parse_args()
    repo = args.repo.resolve()
    source = repo / SOURCE_RELATIVE
    if not source.is_file():
        raise FileNotFoundError(source)
    source_hash = sha256(source)
    if source_hash != EXPECTED_SOURCE_SHA256:
        raise ValueError(f"Publisher workbook SHA-256 mismatch: {source_hash}")

    workbook = load_workbook(source, data_only=True, read_only=True)
    all_records: list[dict[str, Any]] = []
    summaries: dict[str, Any] = {}
    for sheet_name in ("ST1", "ST2"):
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Publisher workbook lacks {sheet_name}")
        records, summary = audit_sheet(workbook[sheet_name], sheet_name)
        all_records.extend(records)
        expected = ARTICLE_REPORTED[sheet_name]
        summary["article_reported"] = expected
        summary["difference_from_article_reported"] = {
            "discovery_rows": summary["discovery_rows"] - expected["discovery_rows"],
            "lead_rows": summary["lead_rows"] - expected["lead_rows"],
            "ukb_available_hits_beta_se_p": summary["all_hit_ukb"]["available_beta_se_p"] - expected["ukb_available_hits"],
            "ukb_nominal_hits_p_lt_0_05": summary["all_hit_ukb"]["nominal_p_lt_0_05"] - expected["ukb_nominal_hits"],
            "ukb_gws_hits_p_lt_5e_8": summary["all_hit_ukb"]["gws_p_lt_5e_8"] - expected["ukb_gws_hits"],
            "ukb_available_leads_beta_se_p": summary["lead_ukb"]["available_beta_se_p"] - expected["ukb_available_leads"],
            "ukb_nominal_leads_p_lt_0_05": summary["lead_ukb"]["nominal_p_lt_0_05"] - expected["ukb_nominal_leads"],
            "ukb_gws_leads_p_lt_5e_8": summary["lead_ukb"]["gws_p_lt_5e_8"] - expected["ukb_gws_leads"],
        }
        summaries[sheet_name] = summary
    workbook.close()

    output_dir = repo / OUTPUT_RELATIVE
    output_dir.mkdir(parents=True, exist_ok=True)
    tsv_path = output_dir / "hfrs_published_lead_variant_replication_2026-09-24.tsv"
    json_path = output_dir / "hfrs_published_lead_variant_replication_2026-09-24.json"
    report_path = output_dir / "hfrs_published_lead_variant_replication_2026-09-24.md"
    with tsv_path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, delimiter="\t", lineterminator="\n")
        writer.writeheader()
        writer.writerows(all_records)

    generated = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    script_path = Path(__file__).resolve()
    payload = {
        "audit": "Descriptive reproduction of publisher-supplement lead-row UKB p-value counts",
        "generated_utc": generated,
        "article": ARTICLE,
        "article_url": "https://www.nature.com/articles/s43587-025-00925-y",
        "source_file": str(SOURCE_RELATIVE),
        "source_sha256": source_hash,
        "script": str(script_path.relative_to(repo)),
        "script_sha256": sha256(script_path),
        "workbook_sheets": ["ST1", "ST2"],
        "availability_definition": "UKB beta, SE and p are all non-missing/non-dot publisher-table cells",
        "significance_definitions": {"nominal": "UKB p < 0.05", "genome_wide": "UKB p < 5e-8"},
        "interpretation_limits": [
            "The publisher tables contain FinnGen discovery hits only, not complete genome-wide summary statistics.",
            "UK Biobank has known cohort-level overlap with many project inputs; exact participant intersections remain unknown.",
            "These source-study variant checks are not sleep-frailty pairwise replication and do not resolve D10 HFRS source/access eligibility.",
            "This audit compares table p-values and availability only; it does not infer cross-cohort effect direction from allele fields.",
            "The public workbook has one fewer available all-hit HFRS row and one fewer nominally significant HFRS lead than the article reports; no undocumented correction is applied.",
        ],
        "results": summaries,
        "detailed_tsv": str(tsv_path.relative_to(repo)),
        "detailed_tsv_sha256": sha256(tsv_path),
    }
    json_path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")

    st1 = summaries["ST1"]
    st2 = summaries["ST2"]
    report = f"""# HFRS published lead-variant replication cross-reference

Generated: {generated}

## Source and method

This descriptive audit reads Supplementary Tables 1 and 2 from the publisher workbook cited by {ARTICLE}. The article says these tables contain FinnGen variants crossing P<5×10⁻⁸ and reports independent leads at r²<0.01; it reports checking those variants in UK Biobank. UKB result availability here requires non-missing beta, SE, and P cells. Nominal significance is P<0.05; genome-wide significance is P<5×10⁻⁸, matching the article's reported cutoffs. The source SHA-256 is `{source_hash}`. Per-variant rows are in `{tsv_path.name}`; machine-readable provenance is in `{json_path.name}`.

## Results

| Table | FinnGen hits | UKB available hits | All-hit P<0.05 / P<5×10⁻⁸ | Independent leads | UKB available leads | Lead P<0.05 / P<5×10⁻⁸ |
|---|---:|---:|---:|---:|---:|---:|
| ST1, HFRS | {st1['discovery_rows']} | {st1['all_hit_ukb']['available_beta_se_p']} (article {st1['article_reported']['ukb_available_hits']}) | {st1['all_hit_ukb']['nominal_p_lt_0_05']} / {st1['all_hit_ukb']['gws_p_lt_5e_8']} (article {st1['article_reported']['ukb_nominal_hits']} / {st1['article_reported']['ukb_gws_hits']}) | {st1['lead_rows']} | {st1['lead_ukb']['available_beta_se_p']} | {st1['lead_ukb']['nominal_p_lt_0_05']} / {st1['lead_ukb']['gws_p_lt_5e_8']} (article {st1['article_reported']['ukb_nominal_leads']} / {st1['article_reported']['ukb_gws_leads']}) |
| ST2, HFRS without dementia | {st2['discovery_rows']} | {st2['all_hit_ukb']['available_beta_se_p']} (article {st2['article_reported']['ukb_available_hits']}) | {st2['all_hit_ukb']['nominal_p_lt_0_05']} / {st2['all_hit_ukb']['gws_p_lt_5e_8']} (article {st2['article_reported']['ukb_nominal_hits']} / {st2['article_reported']['ukb_gws_hits']}) | {st2['lead_rows']} | {st2['lead_ukb']['available_beta_se_p']} | {st2['lead_ukb']['nominal_p_lt_0_05']} / {st2['lead_ukb']['gws_p_lt_5e_8']} (article {st2['article_reported']['ukb_nominal_leads']} / {st2['article_reported']['ukb_gws_leads']}) |

The public ST1 workbook has 1,261/1,588 hits with UKB beta, SE, and P present, compared with 1,262 reported in the article. Its all-hit nominal and genome-wide p-value counts match (688 and 73). Among lead rows, the workbook has 13/36 with P<0.05, compared with 14/36 in the article, and both report 2 with P<5×10⁻⁸. ST2 all-hit and lead availability and p-value counts match the article. These one-row differences are preserved as public-workbook-to-article discrepancies. No p-value was reconstructed or recategorized to force agreement. Article-reported counts are at lines 108–110 of the published Results.

## Interpretation boundary

This reproduces a subset of the Mak et al. source study's own FinnGen-to-UKB lookup. It does not establish independent replication of any sleep–frailty relationship in this project: UKB is shared by several project-side sources, and exact pairwise participant intersections remain unknown. The workbook is hit-conditioned and is not a substitute for full HFRS summary statistics; do not use it for LDSC, LAVA, or genome-wide sleep–HFRS tests. HFRS plan decision D10 remains open pending exact endpoint and authorized source verification. No effect-direction comparison is made because this audit does not independently verify the effect-allele convention for both beta columns.
"""
    report_path.write_text(report, encoding="utf-8")
    print(json.dumps({"ST1": st1, "ST2": st2, "report": str(report_path), "tsv": str(tsv_path)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
