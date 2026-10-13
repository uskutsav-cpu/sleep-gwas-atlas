#!/usr/bin/env python3
"""Metadata-only fifth-snapshot extension; preserve the original four outputs.

Read third-party data in place. Export references and authored classifications,
never the publisher's correlations, standard errors, P values or figures.
"""
import argparse
import csv
import hashlib
import importlib.util
import json
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

EXPECTED = '3488344573bf588d0ad748c83b5a86f5fbd487d711390b6a601661de56587dd7'
DURATION_TRAITS = {'shortsleep', 'longsleep', 'sleepdur', 'accel_sleep_duration'}
PGC_LABELS = {'MDD': 'major depressive disorder', 'BIP': 'bipolar disorder',
              'SCZ': 'schizophrenia', 'OCD': 'obsessive compulsive disorder',
              'ASD': 'autism spectrum disorder', 'ADHD': 'attention deficit hyperactivity disorder'}
PGC_AD_UNRESOLVED_TARGETS = {'alz', 'adhd'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_tsv(path):
    with path.open(newline='', encoding='utf-8-sig') as stream:
        return list(csv.DictReader(stream, delimiter='\t'))


def write_tsv(path, rows, fields):
    with path.open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fields, delimiter='\t', lineterminator='\n')
        writer.writeheader()
        writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--repo', type=Path, required=True)
    parser.add_argument('--sleepchart', type=Path, required=True)
    parser.add_argument('--fg-manifest', type=Path, required=True)
    parser.add_argument('--candidate', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    if digest(args.sleepchart) != EXPECTED:
        raise ValueError('Preserved SleepChart snapshot differs from prior source receipt')
    old_path = args.repo / 'sleep_unified_research_v7/tables/cross_resource_match_annotations.tsv'
    old = read_tsv(old_path)
    groups = defaultdict(list)
    for row in old:
        groups[row['atlas_record_id']].append(row)
    if len(old) != 6548 or len(groups) != 1637 or any(len(x) != 4 for x in groups.values()):
        raise ValueError('Original complete four-snapshot inventory required')
    module_path = args.repo / 'sleep_unified_research_v7/scripts/build_cross_resource_matches.py'
    spec = importlib.util.spec_from_file_location('existing_match_rules', module_path)
    rules = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(rules)
    pheno = {r['trait_id']: r for r in read_tsv(args.repo / 'sleep_unified_research_v4/cell_genomics_readiness_v1/release_candidate/outcome_phenotypes.tsv')}
    fg = {r['phenocode']: r for r in read_tsv(args.fg_manifest)}
    workbook = load_workbook(args.sleepchart, read_only=True, data_only=True)
    sheet = workbook['SF5']
    entries = []
    try:
        for excel_row, values in enumerate(sheet.iter_rows(min_row=4, values_only=True), start=4):
            for block, offset in [('original', 0), ('Dashti', 8), ('Austin', 16)]:
                duration, code, dataset = values[offset:offset + 3]
                if not duration or not code:
                    continue
                # Only presence is used; no estimate/sign/SE/P threshold is selected.
                complete = all(values[offset + i] is not None for i in (3, 4, 6))
                label = fg.get(str(code), {}).get('name', '') if dataset == 'FinnGen' else PGC_LABELS.get(str(code), '')
                entries.append(dict(block=block, offset=offset, row=excel_row,
                                    duration=str(duration), code=str(code), dataset=str(dataset),
                                    label=label, numeric_present=complete,
                                    ref=f'SleepChart2026:SF5:{block}:row{excel_row}:columns{offset+1}-{offset+7}'))
    finally:
        workbook.close()
    block_counts = Counter(e['block'] for e in entries)
    if block_counts != {'original': 1054, 'Dashti': 309, 'Austin': 361}:
        raise ValueError('Unexpected publisher displayed-block inventory')
    original_entries = [e for e in entries if e['block'] == 'original']
    codes = {e['code'] for e in original_entries}
    if len(codes) != 527 or Counter(e['duration'] for e in original_entries) != {'short_vs_normal': 527, 'long_vs_normal': 527}:
        raise ValueError('Original527x2 scope not complete')
    missing_fg = sorted({e['code'] for e in entries if e['dataset'] == 'FinnGen' and not e['label']})
    if missing_fg:
        raise ValueError('FinnGen R9 code metadata unresolved: ' + repr(missing_fg))
    hold = json.loads((args.candidate / 'source_direction_hold.json').read_text())
    # Read own-source hold keys rather than clearing any result by comparator coverage.
    held_keys = {tuple(r['record_key']) for r in hold['affected_core_records']}
    if len(held_keys) != 36:
        raise ValueError('Expected36 exact preserved direction-hold keys')
    annotations, unions = [], []
    for record_id, previous in sorted(groups.items()):
        target = previous[0]
        sleep = target['sleep_trait']
        outcome = target['outcome_trait']
        candidates, ambiguous = [], []
        if sleep in DURATION_TRAITS:
            aliases = rules.aliases_for(pheno[outcome], 'outcome')
            for entry in entries:
                if sleep == 'shortsleep' and 'short' not in entry['duration']:
                    continue
                if sleep == 'longsleep' and 'long' not in entry['duration']:
                    continue
                match = rules.match_type(entry['label'], aliases)
                if match != 'none':
                    candidates.append(entry)
                elif entry['dataset'] == 'PGC' and entry['code'] == 'AD' and outcome in PGC_AD_UNRESOLVED_TARGETS:
                    ambiguous.append(entry)
            populated = [e for e in candidates if e['numeric_present']]
            if populated:
                cls = 'B'
                reason = ('Related duration/outcome result is present. The original UKB tail thresholds/control frame differ; '
                          'displayed legacy blocks do not certify exact outcome release, coding and estimator identity. No exact match is asserted.')
            elif candidates or ambiguous:
                cls = 'E'
                reason = ('A declared outcome slot is blank or its PGC code is ambiguous; no comparison is invented. '
                          'Absence from selected legacy blocks alone cannot establish absent coverage.')
            else:
                cls = 'D'
                reason = ('No counterpart under the fixed existing outcome-alias rules in the complete original527x2 scope. '
                          'This dated table-specific finding does not establish global absence or novelty.')
        else:
            cls = 'D'
            reason = 'No counterpart for this non-duration construct in the declared short/long-duration SF5 scope; no global absence claim.'
        refs = sorted({e['ref'] for e in candidates + ambiguous})
        annotations.append(dict(atlas_record_id=record_id, stage=target['stage'], sleep_trait=sleep,
                                outcome_trait=outcome, comparator='sleepchart_2026', match_class=cls,
                                candidate_result_refs=';'.join(refs), candidate_count=len(refs),
                                source_definition_basis='Primary paper:UKB short<6/long>8 vsnormal6-8;FinnGenR9/PGC;cachedSF5 blocks',
                                adjudication_note=reason))
        before = next(c for c in 'ABECD' if any(r['match_class'] == c for r in previous))
        after = next(c for c in 'ABECD' if c in [r['match_class'] for r in previous] + [cls])
        eligible = (target['stage'] in {'core', 'extension'} and target['h2_z_pass'] == 'True' and target['intercept_pass'] == 'True'
                    and target['sleep_source_id'] not in {'', 'UNKNOWN', 'UNRESOLVED'} and target['outcome_source_id'] not in {'', 'UNKNOWN', 'UNRESOLVED'}
                    and bool(target['sleep_definition'].strip()) and bool(target['outcome_definition'].strip())
                    and target['sleep_source_release'] not in {'', 'UNKNOWN', 'UNRESOLVED'} and target['outcome_source_release'] not in {'', 'UNKNOWN', 'UNRESOLVED'})
        held = (target['stage'], sleep, outcome) in held_keys
        unions.append(dict(atlas_record_id=record_id, stage=target['stage'], before_four=before,
                           added_sleepchart=cls, after_five=after, historical_eligibility=eligible,
                           source_direction_held=held, before_qualified_absent=eligible and before in {'C', 'D'},
                           after_qualified_absent=eligible and after in {'C', 'D'},
                           after_qualified_absent_outside_hold=eligible and after in {'C', 'D'} and not held))
    args.out.mkdir(parents=True, exist_ok=True)
    write_tsv(args.out / 'sleepchart_annotations.tsv', annotations, list(annotations[0]))
    write_tsv(args.out / 'five_snapshot_union.tsv', unions, list(unions[0]))
    summary = dict(original_record_count=len(groups), original_annotation_rows=len(old),
                   new_annotation_rows=len(annotations), combined_annotation_rows=len(old)+len(annotations),
                   source_block_counts=dict(block_counts), original_populated=sum(e['numeric_present'] for e in original_entries),
                   original_blank_slot_refs=[e['ref'] for e in original_entries if not e['numeric_present']],
                   unresolved_PGC_codes=sorted({e['code'] for e in entries if e['dataset']=='PGC' and not e['label']}),
                   sleepchart_class_counts=dict(Counter(r['match_class'] for r in annotations)),
                   before_four_class_counts=dict(Counter(r['before_four'] for r in unions)),
                   after_five_class_counts=dict(Counter(r['after_five'] for r in unions)),
                   transitions=dict(Counter(r['before_four']+'->'+r['after_five'] for r in unions)),
                   qualified_absent_before=sum(r['before_qualified_absent'] for r in unions),
                   qualified_absent_after=sum(r['after_qualified_absent'] for r in unions),
                   qualified_absent_after_outside_hold=sum(r['after_qualified_absent_outside_hold'] for r in unions),
                   before_qualified_by_stage=dict(Counter(r['stage'] for r in unions if r['before_qualified_absent'])),
                   after_qualified_by_stage=dict(Counter(r['stage'] for r in unions if r['after_qualified_absent'])),
                   after_qualified_outside_hold_by_stage=dict(Counter(r['stage'] for r in unions if r['after_qualified_absent_outside_hold'])),
                   source_hashes={'original_four_annotations':digest(old_path),'sleepchart_workbook':digest(args.sleepchart),
                                  'FinnGenR9_manifest':digest(args.fg_manifest),'existing_alias_code':digest(module_path),
                                  'new_script':digest(Path(__file__))},
                   estimator_runs=0, original_files_modified=False,
                   review_status='PENDING_INDEPENDENT_ACTUAL_CLASSIFICATION_REVIEW',
                   scope='Retrospective fifth-snapshot coverage;no global novelty,exact equivalence,biology or user benefit claim')
    (args.out / 'coverage_summary.json').write_text(json.dumps(summary, indent=2)+'\n')
    print(json.dumps({k:summary[k] for k in ['sleepchart_class_counts','after_five_class_counts','qualified_absent_before','qualified_absent_after','qualified_absent_after_outside_hold']}))


if __name__ == '__main__':
    main()
