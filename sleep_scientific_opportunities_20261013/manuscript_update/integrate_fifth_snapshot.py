"""Update the preserved integrated manuscript with completed evidence only."""
from copy import deepcopy
from pathlib import Path
import argparse
import hashlib
import json
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SOURCE = Path('/Users/swethasunilkumar/Documents/Codex/2026-10-08/go/outputs/Cell_Genomics_Scientific_Validation_2026-10-12/Integrated_Sleep_Atlas_With_Completed_AI_Feasibility.docx')
ROOT = Path(__file__).resolve().parents[1]

def insert_before(anchor, text, style):
    return anchor.insert_paragraph_before(text, style)

def main(outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    original_sha = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    doc = Document(SOURCE)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace('12 October 2026', '13 October 2026')
    paragraphs = list(doc.paragraphs)
    assert 'four accessible row-level snapshots' in paragraphs[20].text
    summary = json.loads((ROOT/'resource_coverage/results_corrected_v3/coverage_summary.json').read_text())
    assert summary['after_five_class_counts'] == {'B':1193, 'D':362, 'E':82}
    paragraphs[5].text = ('Interpreting genetic relationships between sleep and disease requires attention to phenotype definitions, source releases and uncertainty. We assembled a source-aware resource from 12 sleep phenotypes, 396 core correlations and 1,200 extension correlations, preserving separate testing families. Native reproduction agreed with historical logs at printed precision. All 41 estimated external checks reused discovery sleep data and represent outcome-side validation. Comparison with five resource snapshots yielded 8,185 annotations without certifying exact source-and-estimand equivalence. Primary-paper controls challenged assumed bowel-source effect directions, prompting holds on 36 core records. Unauthenticated MVP uncertainty prompted an additional inferential hold on five historical validation records and one exploratory fit. Exact-key retrieval integrates source definitions, original testing families, QC and interpretation holds. A ten-task computational feasibility subset found a matched-evidence accuracy ceiling in both representations. Improved researcher interpretation and efficiency remain unmeasured.')
    paragraphs[8].text += ' SleepChart adds categorical duration and disease-genetics context [15].'
    paragraphs[20].text = ('We compared each of 1,637 historical result records with five row-level snapshots, yielding 8,185 annotations (Table 2). The fixed population comprises 396 core, 1,200 extension and 41 estimated validation records; it does not include all 217 external candidates. No record was certified as an exact source-and-estimand match. Across the union, 1,193 records had related but nonidentical counterparts, 362 were absent within the mapped snapshot scope and 82 had insufficient comparability evidence. Adding the cached SleepChart SF5 table changed ten union classifications: four core colorectal-cancer duration records and six estimated validation records acquired related prior context. These are coverage changes, not new associations (Figure 2). The original four annotations were preserved and reused; their phenotype semantics were not comprehensively re-audited in this extension.')
    paragraphs[21].text = ('Among historically source/QC-qualified core and extension records, 357 had no exact, related or unresolved counterpart within the five snapshots and declared mapping: 81 core and 276 extension records. Three core records remain under a bowel-source direction hold, leaving 354 outside that hold (78 core and 276 extension). These dated, vocabulary-bounded counts do not establish global literature absence, statistical significance or first discovery. The original four-snapshot analysis remains preserved separately. A direct CTG-VL export was not included. The new SleepChart annotation has 110 related, eight incompatible and 1,519 within-scope absent assignments; categorical duration tails and qualified disease proxies do not establish exact phenotype or source equivalence. Our initial melanoma and BMI-composite matching errors were corrected before independent acceptance, with prior outputs retained. All 36 direction holds remain operative.')
    paragraphs[22].text = 'Table 2  Matching within five resource snapshots'
    paragraphs[24].text = ('Each snapshot is aligned to the same 1,637 historical records. Incompatible marks demonstrable construct differences, including explicitly excluded or composite outcomes. Absent denotes the declared mapped snapshot scope; unresolved denotes insufficient identity evidence. The union follows the frozen precedence rule and is not a sum of columns. SleepChart presence does not clear source-direction or MVP uncertainty holds.')
    paragraphs[39].text = ('The comparator benchmark situates the resource within extensive prior sleep-genetics work. Factor, composite, organ-focused and categorical-duration analyses already exist [3-5,15]. Adding SleepChart reduced the records with mapped absence outside direction holds from 358 to 354; it did not establish a new genetic relationship. The extension also illustrates why label-only matching can mistake an excluded diagnosis or composite outcome for the requested phenotype. Snapshot-bounded gaps may identify complementary resource configurations, while inherited earlier annotations and a fixed mapping vocabulary limit completeness. The direction assessment was purposive: 22 conflicting controls cannot estimate a representative error prevalence across genomic resources.')
    paragraphs[49].text = ('Authored code, scientific reports and reviewed comparator classifications are maintained at https://github.com/uskutsav-cpu/sleep-gwas-atlas on research/scientific-opportunities-2026-10-13, building on the preserved 12 October validation branch. The 13 October increment starts at commit 053e31dacb4a0012d64d1a5e3d551f9da3145b2c; later commits preserve source-acquisition and preprocessing preparation separately. This is not an archival deposition. Raw GWAS bodies, publisher supplements, private scoring keys and protected comparator extracts are not redistributed through the branch. The machine-readable candidate remains available locally for controlled review subject to source terms. Permitted redistribution, archival access and investigator approval remain required for public release.')
    paragraphs[67].text += (' The retrospective fifth-snapshot extension used the exact cached SleepChart SF5 workbook (SHA256 3488344573bf588d0ad748c83b5a86f5fbd487d711390b6a601661de56587dd7) and official FinnGen R9 endpoint metadata [15]. Original SF5 contains 527 endpoints for each of two duration tails, with four blank slots; the Dashti and Austin blocks are selected legacy displays. Ambiguous codes and incompatible candidates were separated from authenticated related context. Independent reconstruction reproduced every new classification and union count. The 9 October publisher correction concerns Figure 1 axis labels [16]; it does not announce an SF5 replacement. The cached workbook was not certified byte-identical to a current publisher binary.')
    paragraphs[70].text = paragraphs[70].text.replace('its completed results and limitations belong in a separate validation addendum.', 'the completed ten-task computational subset is reported above and in the full validation addendum. Its AI evidence is separate from the unexecuted human pilot.')
    paragraphs[110].text = 'Table S11  Preserved four-snapshot comparator summary and union. File: cross_resource_match_summary.tsv.'

    # Keep the article coherent: completed AI evidence belongs in Results/Methods.
    anchor = paragraphs[37]
    insert_before(anchor, 'AI feasibility shows a matched evidence ceiling', 'Heading 2')
    insert_before(anchor, 'The prospectively selected ten-task subset produced 90 fresh condition evaluations using gpt-6.1-sol, gpt-6-sol and gpt-6-luna. These are three OpenAI aliases within one provider family; backend versions were not independently authenticated. Two conventional-condition cells returned incorrect packet digests and were excluded under the frozen rule, leaving 88 included cells. The matched conventional presentation and atlas presentation each achieved 30/30 exact-gold agreement, with zero paired differences. Conventional-resource exact agreement was 4/10, 7/10 and 5/8 for the three aliases, respectively. All reported outcomes, exclusions and original responses are retained in the addendum.', 'Normal')
    insert_before(anchor, 'The conventional condition lacked exact target evidence; its strict-score errors mainly encoded conservative nulls where the gold specified false or NOT_ESTABLISHED. The reviewed responses did not substitute unsupported target associations. Thus condition differences cannot identify an interface benefit when evidence availability differs. The information-matched comparison showed a ceiling on this finite exact-record retrieval task. It establishes neither improved scientific interpretation nor equivalence. Source rules and verdict fields were supplied, questions and aliases are dependent, access isolation was instruction-based, and independently observed read traces and reliable cost/timing measures were unavailable. No human benefit was measured.', 'Normal')
    methods_anchor = paragraphs[69]
    insert_before(methods_anchor, 'Computational feasibility evaluation', 'Heading 2')
    insert_before(methods_anchor, 'We retained the separate 60-final-task and six-development-task benchmark and its frozen rubric. Before responses, ten final tasks were selected by an outcome-independent seeded schedule, with a fresh cell for each task, alias and condition. The three conditions were conventional-resource evidence, the same factual evidence in conventional tables, and the atlas presentation. Available model aliases were run at medium reasoning effort through Codex; deterministic sampling and backend-version pinning were not authenticated. Full responses, failures and attempted retries were recorded. Frozen exact-schema scoring was independently reviewed by AI assistants and remains distinct from accountable human semantic assessment. Paired B/C counts are descriptive for the evaluated questions and aliases; no provider-population or equivalence inference is made. The proposed 540-cell expansion was not executed.', 'Normal')

    table = doc.tables[1]
    table.add_column(Inches(.8))
    table.add_row()
    values = [
        ['Snapshot','Exact','Related','Incompatible','Absent','Unresolved'],
        ['Human GWAS ATLAS release 3','0','210','0','1,344','83'],
        ['Morrison 2024 S10-S15','0','888','0','749','0'],
        ['Goodman 2025 SD24','0','492','0','815','330'],
        ['Fan 2026 SD2 SD4 SD11','0','40','0','1,549','48'],
        ['SleepChart 2026 cached SF5','0','110','8','1,519','0'],
        ['Union across five snapshots','0','1,193','0','362','82']]
    widths = [2.7,.55,.75,1.05,.8,1.05]
    table.autofit = False
    for col, width in zip(table.columns, widths): col.width = Inches(width)
    for row_index, (row, texts) in enumerate(zip(table.rows, values)):
        for col_index, (cell, text, width) in enumerate(zip(row.cells, texts, widths)):
            cell.width = Inches(width); cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            cell.paragraphs[0].paragraph_format.space_before = Pt(3)
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(9); run.font.bold = row_index == 0
                run.font.color.rgb = RGBColor.from_string('FFFFFF' if row_index == 0 else '000000')
            props = cell._tc.get_or_add_tcPr()
            for old in props.findall(qn('w:shd')):props.remove(old)
            shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'263C50' if row_index == 0 else ('F1F4F6' if row_index%2 == 0 else 'FFFFFF'));props.append(shading)
            for old in props.findall(qn('w:tcBorders')):props.remove(old)
            borders=OxmlElement('w:tcBorders');props.append(borders)
            for edge in ['top','left','bottom','right']:
                border=OxmlElement('w:'+edge);border.set(qn('w:val'),'single');border.set(qn('w:sz'),'4');border.set(qn('w:color'),'D9D9D9');borders.append(border)
            for old in props.findall(qn('w:tcMar')):props.remove(old)
            margin=OxmlElement('w:tcMar');props.append(margin)
            for edge in ['top','left','bottom','right']:
                val=OxmlElement('w:'+edge);val.set(qn('w:w'),'90');val.set(qn('w:type'),'dxa');margin.append(val)
    if table.rows[0]._tr.trPr is None:table.rows[0]._tr.get_or_add_trPr()
    repeat=OxmlElement('w:tblHeader');table.rows[0]._tr.trPr.append(repeat)

    refs_anchor = paragraphs[88]
    p = insert_before(refs_anchor, '15. The MULTI Consortium et al. Sleep chart of biological ageing clocks in middle and late life. Nature (2026). doi:10.1038/s41586-026-10524-5. https://www.nature.com/articles/s41586-026-10524-5', 'Normal')
    p.paragraph_format.space_after = Pt(5)
    p = insert_before(refs_anchor, '16. The MULTI Consortium et al. Publisher Correction to Sleep chart of biological ageing clocks in middle and late life. Nature (9 October 2026). doi:10.1038/s41586-026-11273-1. https://www.nature.com/articles/s41586-026-11273-1', 'Normal')
    p.paragraph_format.space_after = Pt(5)
    figure_anchor=paragraphs[93]
    figure_anchor.paragraph_format.page_break_before=True
    heading=insert_before(figure_anchor,'Figure 2  Dated comparator coverage','Heading 1')
    heading.paragraph_format.page_break_before=True
    image=insert_before(figure_anchor,'','Normal')
    image.add_run().add_picture(str(ROOT/'resource_coverage/figures/five_snapshot_coverage_increment.png'),width=Inches(6.7))
    image.paragraph_format.keep_with_next=True
    caption_text=(ROOT/'resource_coverage/figures/CAPTION.txt').read_text().strip()
    caption_text=caption_text.replace('fixed1,637','fixed 1,637').replace('396core,1,200extension and41estimated','396 core, 1,200 extension and 41 estimated')
    caption=insert_before(figure_anchor,caption_text,'Caption')
    caption.paragraph_format.keep_together=True
    doc.add_paragraph('Table S14  Authored SleepChart snapshot annotations and candidate references. File: sleepchart_annotations.tsv. Current accepted version: results_corrected_v3; no publisher genetic estimates are redistributed.','Normal')
    doc.add_paragraph('Table S15  Five-snapshot union and hold-qualified count dispositions. File: five_snapshot_union.tsv. All records and original testing families remain retained.','Normal')
    for paragraph in doc.paragraphs:
        revised = paragraph.text.replace('before independent acceptance', 'before independent AI review')
        revised = revised.replace('three OpenAI aliases within one provider family', 'three model aliases from one provider (OpenAI)')
        revised = revised.replace('The three conditions were conventional-resource evidence, the same factual evidence in conventional tables, and the atlas presentation.', 'Condition A supplied conventional-resource evidence separately. Conditions B and C shared identical factual evidence, presented in conventional tables and the atlas, respectively.')
        if revised != paragraph.text:
            paragraph.text = revised
    target=outdir/'Integrated_Sleep_Atlas_With_Five_Resource_Comparison.docx'
    doc.save(target)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==original_sha
    text='\n'.join(p.text for p in doc.paragraphs)+'\n'
    (outdir/'Integrated_Manuscript_Text.txt').write_text(text)
    (outdir/'MANUSCRIPT_UPDATE_RECEIPT.json').write_text(json.dumps({'source':str(SOURCE),'source_sha256':original_sha,'original_preserved':True,'output':str(target),'output_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),'changed_scope':'Accepted fifth-snapshot coverage and completed finite AI results; new source fits/human utility remain pending.','new_genetic_estimates':0,'human_responses':0,'render_status':'PENDING','figure_source':'resource_coverage/figures/five_snapshot_coverage_increment.png'},indent=2)+'\n')
    print(target)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--out',type=Path,required=True)
    main(parser.parse_args().out)
