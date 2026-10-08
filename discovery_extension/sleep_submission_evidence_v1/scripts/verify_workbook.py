#!/usr/bin/env python3
"""Read-only validation of a static research workbook export, not Excel recalculation."""
import argparse, collections, csv, hashlib, json, math, platform
from pathlib import Path
import openpyxl

P=Path(__file__).resolve().parents[1]

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def close_float(a,b):
    return math.isfinite(a) and math.isfinite(b) and abs(a-b)<=8*max(math.ulp(a),math.ulp(b))

def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--intermediate',type=Path,required=True)
    args=parser.parse_args()
    book=P/'tables/supplementary_data.xlsx'
    data=json.loads(args.intermediate.read_text())
    original=digest(book); issues=[]; sheets=[]; missing_counts=collections.Counter(); typed_blank_strings=0
    files={str(book):original,str(args.intermediate):digest(args.intermediate)}
    w=openpyxl.load_workbook(book,read_only=True,data_only=False)
    expected_names=['Overview']+[x['sheet'] for x in data]
    if len(data)!=8 or w.sheetnames!=expected_names:issues.append({'kind':'sheet_names','actual':w.sheetnames,'expected':expected_names})
    for table in data:
        name=table['sheet']; src=P/table['source'];files[str(src)]=digest(src)
        with src.open(newline='') as f:
            raw=list(csv.reader(f,delimiter='\t'))
        headers=table['headers']; n=len(headers); sh=w[name]
        if not raw or raw[0]!=headers:issues.append({'kind':'source_header','sheet':name})
        if len(raw)-1!=len(table['rows']):issues.append({'kind':'source_rows','sheet':name,'actual':len(raw)-1,'expected':len(table['rows'])})
        all_actual=list(sh.iter_rows())
        expected=[headers]+table['rows']
        types=collections.Counter();checked=0;max_ulp=0.0;markers=collections.Counter()
        for i,row in enumerate(expected):
            for j,value in enumerate(row):
                checked+=1
                cell=all_actual[i][j] if i<len(all_actual) and j<len(all_actual[i]) else None
                actual=cell.value if cell is not None else None
                dtype=cell.data_type if cell is not None else 'missing'
                types[dtype]+=1
                spot={'sheet':name,'row':i+1,'column':j+1,'header':headers[j]}
                if i and i<len(raw) and j<len(raw[i]):
                    text=raw[i][j]
                    if isinstance(value,bool): same=text.lower() in ['true','false'] and value==(text.lower()=='true')
                    elif isinstance(value,(int,float)):
                        try:same=close_float(float(text),float(value))
                        except ValueError:same=False
                    else:same=text==value
                    if not same:issues.append(dict(spot,kind='TSV_JSON_VALUE',source_value=text,intermediate_value=value))
                if isinstance(value,bool):good=dtype=='b' and type(actual) is bool and actual==value
                elif isinstance(value,(int,float)):
                    good=dtype=='n' and type(actual) in (float,int) and close_float(float(value),float(actual))
                    if good:max_ulp=max(max_ulp,abs(float(value)-float(actual))/max(math.ulp(float(value)),math.ulp(float(actual))))
                elif isinstance(value,str):
                    good=dtype in ['s','inlineStr'] and isinstance(actual,str) and actual==value
                    # Explicit OOXML t="str" empty cells preserve a TSV empty field.
                    # Missing NA/UNRESOLVED labels never receive this exception.
                    if value=='' and actual is None and dtype=='str':
                        good=True;typed_blank_strings+=1
                    if value in ['NA','NaN','nan','NAN','N/A','','NOT_SEARCHED','UNRESOLVED','NOT_AVAILABLE']:
                        markers[value]+=1; missing_counts[value]+=1
                else:good=value is None and actual is None
                if not good:issues.append(dict(spot,kind='XLSX_VALUE_OR_CELLTYPE',expected_value=value,actual_value=actual,actual_type=dtype))
        for i,row in enumerate(all_actual):
            for j,cell in enumerate(row):
                if (i>=len(expected) or j>=n) and cell.value is not None:
                    issues.append({'sheet':name,'row':i+1,'column':j+1,'kind':'EXTRA_NONEMPTY_CELL'})
        sheets.append({'sheet':name,'source':table['source'],'data_rows':len(table['rows']),'columns':n,'checked_cells_including_headers':checked,'cell_types':dict(types),'preserved_missing_marker_counts':dict(markers),'maximum_numeric_roundtrip_ulps':max_ulp})
    # Cached values are checked independently without running an Excel engine.
    wc=openpyxl.load_workbook(book,read_only=True,data_only=True)
    expected_counts=[1200,603,217,23,0,7,0]
    counts=[wc['Overview'].cell(row=r,column=2).value for r in range(5,12)]
    if counts!=expected_counts:issues.append({'kind':'OVERVIEW_CACHED_COUNTS','actual':counts,'expected':expected_counts})
    formula_cells={f'B{r}':w['Overview'].cell(row=r,column=2).value for r in [5,6,7,8,10]}
    if any(not isinstance(v,str) or not v.startswith('=') for v in formula_cells.values()):issues.append({'kind':'OVERVIEW_FORMULAS_MISSING'})
    nov=next(t for t in data if t['sheet']=='Novelty'); idx={h:i for i,h in enumerate(nov['headers'])}
    classes=collections.Counter(r[idx['current_novelty_class']] for r in nov['rows'])
    direct='PREVIOUSLY_PUBLISHED_SUBSTANTIALLY_SIMILAR'
    repl_direct=sum(r[idx['current_novelty_class']]==direct and r[idx['historical_replication_class']]=='REPLICATED' for r in nov['rows'])
    if classes[direct]!=83 or repl_direct!=15:issues.append({'kind':'STALE_NOVELTY_CLASSIFICATION','direct':classes[direct],'replicated_direct':repl_direct})
    with (P/'tables/NOVELTY_REPLICATED_23.tsv').open(newline='') as f: nr=list(csv.DictReader(f,delimiter='\t'))
    novel_by_pair={r[idx['pair_id']]:dict(zip(nov['headers'],r)) for r in nov['rows']}
    for row in nr:
        observed=novel_by_pair.get(row['pair_id'])
        if observed is None or observed['current_novelty_class']!=row['current_novelty_class']:issues.append({'kind':'NOVELTY_REPLICATION_JOIN','pair_id':row['pair_id']})
    w.close();wc.close()
    final=digest(book)
    if original!=final:issues.append({'kind':'WORKBOOK_CHANGED_DURING_VALIDATION'})
    receipt={'date':'2026-10-08','status':'PASS' if not issues else 'FAIL','validation_scope':'ALL_EIGHT_TABS_ALL_VALUES_AND_CELL_TYPES_AGAINST_CANONICAL_TSV_AND_TYPED_INTERMEDIATE;CACHED_OVERVIEW_ONLY','native_Excel_recalculation_performed':False,'workbook_writes':False,'floating_tolerance':'8 IEEE754 ulps;no fixed absolute floor;all nonempty strings/booleans exact','empty_string_serialization':'Explicit OOXML t=str blank cells read as None by openpyxl;semantically empty,never converted to zero or a missing label','empty_strings_exported_as_typed_blank_cells':typed_blank_strings,'python_version':platform.python_version(),'openpyxl_version':openpyxl.__version__,'sheet_names':expected_names,'table_sheets':sheets,'total_checked_table_cells':sum(x['checked_cells_including_headers'] for x in sheets),'total_data_rows':sum(x['data_rows'] for x in sheets),'overview_cached_counts':counts,'overview_formula_cells':formula_cells,'novelty_classes':dict(classes),'novelty_direct_prior_comparators':classes[direct],'replicated_direct_prior_comparators':repl_direct,'preserved_missing_marker_counts':dict(missing_counts),'source_hashes':files,'workbook_hash_unchanged_during_check':original==final,'issues':issues}
    (P/'logs/workbook_independent_validation.json').write_text(json.dumps(receipt,indent=2,allow_nan=False)+'\n')
    print(json.dumps({'status':receipt['status'],'checked_cells':receipt['total_checked_table_cells'],'data_rows':receipt['total_data_rows'],'overview_cached_counts':counts,'issues':len(issues)}))
    if issues:raise SystemExit(1)

if __name__=='__main__':main()
