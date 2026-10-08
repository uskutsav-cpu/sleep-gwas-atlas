#!/usr/bin/env python3
"""Six editable numerical table files, no manuscript prose."""
from pathlib import Path
from datetime import datetime,timezone
import csv, json, zipfile, io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
P=Path(__file__).resolve().parents[1]
OUT=P/'tables/word';OUT.mkdir(exist_ok=True)

def read(file):
    with (P/'tables'/file).open(newline='') as f:return list(csv.DictReader(f,delimiter='\t'))
def num(v):
    return 'NA' if v=='NA' else f'{float(v):.4f}'
def prob(v):
    return 'NA' if v=='NA' else f'{float(v):.3g}'
def make(name,title,headers,rows,widths,note,source):
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.65);s.left_margin=s.right_margin=Inches(.7)
    for sty in ['Normal','Title','Heading 1']:
        doc.styles[sty].font.name='Arial';doc.styles[sty].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10);doc.styles['Title'].font.size=Pt(14)
    for border in list(doc.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    doc.add_paragraph(title,'Title')
    tbl=doc.add_table(rows=1,cols=len(headers));tbl.alignment=WD_TABLE_ALIGNMENT.CENTER;tbl.autofit=False
    for c,w in zip(tbl.columns,widths):c.width=Inches(w)
    for cell,h,w in zip(tbl.rows[0].cells,headers,widths):cell.text=h;cell.width=Inches(w)
    hdr=OxmlElement('w:tblHeader');tbl.rows[0]._tr.get_or_add_trPr().append(hdr)
    for vals in rows:
        cells=tbl.add_row().cells
        for cell,value,w in zip(cells,vals,widths):cell.text=str(value);cell.width=Inches(w)
    for i,row in enumerate(tbl.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                m=OxmlElement('w:'+side);m.set(qn('w:w'),'70');m.set(qn('w:type'),'dxa');margins.append(m)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'24445D' if i==0 else ('F1F4F7' if i%2==0 else 'FFFFFF'));pr.append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.space_before=Pt(0);para.paragraph_format.line_spacing=1.05
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    run.font.size=Pt(9);run.font.bold=i==0;run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    doc.add_paragraph(note)
    doc.add_paragraph('Numerical source: '+source)
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.created=doc.core_properties.modified=datetime(2026,10,8,tzinfo=timezone.utc)
    raw=io.BytesIO();doc.save(raw)
    # Normalize zip timestamps so repeated exports are deterministic.
    with zipfile.ZipFile(raw) as src,zipfile.ZipFile(OUT/name,'w',compression=zipfile.ZIP_DEFLATED) as target:
        for entry in src.infolist():
            info=zipfile.ZipInfo(entry.filename,(2026,10,8,0,0,0));info.compress_type=zipfile.ZIP_DEFLATED;target.writestr(info,src.read(entry.filename))
    return {'file':'tables/word/'+name,'data_rows':len(rows),'columns':len(headers),'source':source}

def main():
    receipts=[]
    x=read('main_panel_composition.tsv')
    receipts.append(make('table1_panel_composition.docx','Table 1 Panel composition',
        ['Analysis','Sleep traits','External traits','Planned pairs'],[[r['analysis'].replace('_',' '),r['sleep_traits'],r['non_sleep_traits'],r['planned_tests']] for r in x],[2.7,1.3,1.3,1.5],
        'The locked core remains historical and distinct from the exploratory extension. Full original core checkpoint verification remains incomplete.','main_panel_composition.tsv'))
    x=read('phenotype_and_source_metadata.tsv');selected=[r for r in x if r['role']!='DISCOVERY_OUTCOME']
    # Preserve the entire 125-row detailed source metadata in TSV/XLSX; main table groups the Pan-UKB panel.
    if len(selected)==len(x):selected=[r for r in x if 'panukbb' not in r['trait_id']]
    rows=[[r['trait_id'].removeprefix('FINNGEN_R13_').replace('_',' '),r['release'].replace('_',' '),r['ancestry'],r['sample_size'],r['cases'],r['controls']] for r in selected]
    pan=[r for r in x if 'panukbb' in r['trait_id']]
    if pan:
        ns=[int(float(r['sample_size'])) for r in pan if r['sample_size']!='NA']
        rows.append([f'Pan UKB selected panel ({len(pan)} phenotypes)',pan[0]['release'],'EUR',f'{min(ns):,} to {max(ns):,}','Varies','Varies'])
    receipts.append(make('table2_source_characteristics.docx','Table 2 GWAS source characteristics',
        ['Trait or source','Release','Ancestry','N','Cases','Controls'],rows,[1.95,1.45,.75,.75,.95,.95],
        'Exact definitions, models, effective sample sizes, URLs and source identities for all 125 source records are retained in the detailed metadata. NA is unavailable or inapplicable, never zero.','phenotype_and_source_metadata.tsv'))
    x=read('main_discovery_summary.tsv')
    selected_keys=['candidate_phenotypes_considered','discovery_tests','external_traits','external_h2_primary_pass','sleep_traits','bh_significant','replication_family','replication_tests','replicated_external_phenotypes',
                   'replicated_exact','replicated_comparable','source_level_native_ldsc_reruns','replication_external_outcome_side_replication',
                   'replication_directionally_concordant','replication_underpowered','replication_no_independent_dataset']
    rows=[[r['metric'].replace('_',' '),r['value']] for r in x if r['metric'] in selected_keys]
    receipts.append(make('table3_outcome_summary.docx','Table 3 Discovery and validation accounting',['Measure','Count'],rows,[5.3,1.5],
        'Discovery BH correction uses 1,200 pairs. Validation Bonferroni uses all 217 selected candidates, including unavailable and ineligible outcomes. Native replay remains blocked.','main_discovery_summary.tsv'))
    x=read('replicated_23.tsv');rows=[]
    for r in x:
        rows.append([r['sleep_trait'].replace('_',' ')+' / '+r['external_phenotype_name'],num(r['discovery_rg'])+' ('+num(r['discovery_se'])+')',
                     num(r['replication_rg'])+' ('+num(r['replication_se'])+')',prob(r['replication_p']),
                     'C' if 'COMPARABLE' in r['phenotype_match_status'] else 'E',prob(r['heterogeneity_p'])])
    receipts.append(make('table4_outcome_side_pairs.docx','Table 4 Outcome side validated pairs',
        ['Sleep and external phenotype','Discovery rg (SE)','Validation rg (SE)','Validation P','Match','Nominal heterogeneity P'],rows,[2.2,1.0,1.0,.8,.7,1.1],
        'Match: E = exact definition; C = comparable definition. All 23 reuse the sleep GWAS. None is fully independent two-trait replication. Threshold: 0.05/217. Heterogeneity assumes zero estimator covariance; seven nominal flags are not a family-corrected count. Archived-precision qualifications and 95% intervals remain in TSV.','replicated_23.tsv'))
    x=read('main_phenotype_domains.tsv')
    receipts.append(make('table5_phenotype_domains.docx','Table 5 Phenotype domain results',['Domain','Phenotypes','Comparisons','BH positive pairs'],
        [[r['phenotype_domain'].replace('_',' '),r['distinct_phenotypes'],r['comparisons'],r['bh_significant']] for r in x],[3.2,1.1,1.2,1.3],
        'Counts describe correlated comparisons. They are not independent diseases, biological discoveries or loci.','main_phenotype_domains.tsv'))
    x=read('heterogeneity_7.tsv')
    rows=[[r['sleep_trait'].replace('_',' ')+' / '+r['external_phenotype_name'],num(r['discovery_rg'])+' ('+num(r['discovery_se'])+')',
           num(r['replication_rg'])+' ('+num(r['replication_se'])+')',num(r['difference_replication_minus_discovery']),prob(r['heterogeneity_p'])] for r in x]
    receipts.append(make('table6_nominal_heterogeneity.docx','Table 6 Nominally heterogeneous validated pairs',
        ['Sleep and external phenotype','Discovery rg (SE)','Validation rg (SE)','Validation minus discovery','Nominal P'],rows,[2.6,1.1,1.1,1.1,.9],
        'Historical P<0.05 selection preserved. Shared sleep input implies unknown estimator covariance. Full-family correction and covariance-grid results are retrospective sensitivities, not new confirmatory evidence.','heterogeneity_7.tsv'))
    (P/'logs/word_tables_receipt.json').write_text(json.dumps(receipts,indent=2)+'\n')
    print(json.dumps(receipts))
if __name__=='__main__':main()
